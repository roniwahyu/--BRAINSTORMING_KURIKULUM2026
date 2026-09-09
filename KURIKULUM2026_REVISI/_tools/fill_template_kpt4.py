# -*- coding: utf-8 -*-
"""Final cleanup: fix remaining text placeholders, clear leftover rows, fill matriks CPL-MK (T29)."""
import copy
from docx import Document

SRC = "D:/!!MYDOCUMENTS2026/!!!SISTEKIN2026/!!BRAINSTORMING_KURIKULUM2026/KURIKULUM2026_REVISI/BUKU_KPT_SISTEKIN_2026_FINAL.docx"
doc = Document(SRC)

def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold

# --- Global text replacement (remaining minor placeholders) ---
REPL = {
    "[Email Prodi]": "[PERLU DATA UPPS]",
    "[bidang keilmuan utama prodi]": "sistem dan teknologi informasi cerdas",
    "[Nama MK 1]": "Pengantar Sistem dan Teknologi Informasi",
    "[Nama MK 2]": "Analisis dan Perancangan Sistem Informasi",
    "[Nama MK 3]": "Machine Learning",
    "[Nama MK 4]": "Komputasi Awan (Cloud Computing)",
    "[Nama MK 5]": "Integrasi Layanan Cerdas Berbasis AI",
    "[Nama MK 6]": "UI/UX Design & Prototyping",
    "[Nama MK 7]": "Inovasi Teknologi dan Startup Digital",
    "[Nama MK 8]": "Capstone Project FSTI",
}
for p in doc.paragraphs:
    for run in p.runs:
        for k, v in REPL.items():
            if k in run.text:
                run.text = run.text.replace(k, v)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    for k, v in REPL.items():
                        if k in run.text:
                            run.text = run.text.replace(k, v)

# --- Table 19: clear leftover rows 6,7 (already filled 5 meaningful rows) ---
t = doc.tables[19]
for i in range(6, len(t.rows)):
    for c in t.rows[i].cells:
        set_cell(c, "")

# --- Table 29: Matriks CPL-MK (Mata Kuliah x 9 CPL cols) ---
# Fill 8 representative MKs with their CPL support (simplified I-R-M)
t = doc.tables[29]
mk_cpl = [
    ("Pengantar Sistem dan Teknologi Informasi", ["","","","","√","√","","",""]),
    ("Analisis dan Perancangan Sistem Informasi", ["","","","","","√","","√",""]),
    ("Machine Learning", ["","","","","√","","","","√"]),
    ("Komputasi Awan (Cloud Computing)", ["","","","","","","√","",""]),
    ("Integrasi Layanan Cerdas Berbasis AI", ["","","","","","","","","√"]),
    ("UI/UX Design & Prototyping", ["","","","","","","","√",""]),
    ("Inovasi Teknologi dan Startup Digital", ["√","√","√","√","","","","","√"]),
    ("Capstone Project FSTI", ["√","√","√","√","√","√","√","√","√"]),
]
for i in range(1, len(t.rows)):
    cells = t.rows[i].cells
    if i - 1 < len(mk_cpl):
        set_cell(cells[0], mk_cpl[i-1][0])
        for j in range(1, min(len(cells), 10)):
            set_cell(cells[j], mk_cpl[i-1][1][j-1] if j-1 < len(mk_cpl[i-1][1]) else "")
    else:
        for c in cells:
            set_cell(c, "")

doc.save(SRC)
print("SAVED. Final cleanup done.")
