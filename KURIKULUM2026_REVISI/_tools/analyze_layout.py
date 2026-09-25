import fitz
import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
base_dir = r"d:\!!MYDOCUMENTS2026\!!!SISTEKIN2026\!!BRAINSTORMING_KURIKULUM2026\KURIKULUM2026_REVISI"

def analyze_pdf_tables_and_metadata(pdf_name):
    filepath = os.path.join(base_dir, pdf_name)
    doc = fitz.open(filepath)
    print("=" * 80)
    print(f"ANALYSIS OF: {pdf_name}")
    print(f"Total Pages: {len(doc)}")
    print("=" * 80)
    
    for page_num in range(len(doc)):
        page = doc[page_num]
        print(f"\n>>> PAGE {page_num + 1} (Size: {page.rect.width} x {page.rect.height}) <<<")
        # Check orientation: Portrait vs Landscape
        orientation = "Landscape" if page.rect.width > page.rect.height else "Portrait"
        print(f"Orientation: {orientation}")
        
        # Check drawings / table lines / rects
        drawings = page.get_drawings()
        print(f"Drawing paths / rects count: {len(drawings)}")
        
        # Extract blocks
        blocks = page.get_text("blocks")
        for b in blocks:
            # b: (x0, y0, x1, y1, text, block_no, block_type)
            txt = b[4].strip().replace("\n", " | ")
            if txt:
                print(f"  [{b[0]:.1f}, {b[1]:.1f}] {txt[:120]}")

analyze_pdf_tables_and_metadata("0001_Template_RPS_OBE.pdf")
print("\n" * 3)
analyze_pdf_tables_and_metadata("Template_RPS_GEN_2026_20092026-SWI.pdf")

# Also check docx if available
docx_path = os.path.join(base_dir, "Template_RPS_GEN_2026_20092026-SWO.docx")
if os.path.exists(docx_path):
    print("\n" * 3)
    print("=" * 80)
    print("ANALYSIS OF DOCX: Template_RPS_GEN_2026_20092026-SWO.docx")
    print("=" * 80)
    doc = docx.Document(docx_path)
    print(f"Paragraphs count: {len(doc.paragraphs)}")
    print(f"Tables count: {len(doc.tables)}")
    for t_idx, tbl in enumerate(doc.tables):
        print(f"\n--- Table {t_idx + 1} ({len(tbl.rows)} rows x {len(tbl.columns)} cols) ---")
        for r_idx in range(min(5, len(tbl.rows))):
            row_vals = [c.text.strip().replace("\n", " ") for c in tbl.rows[r_idx].cells]
            print(f"  Row {r_idx + 1}: {' | '.join(row_vals[:6])}")
