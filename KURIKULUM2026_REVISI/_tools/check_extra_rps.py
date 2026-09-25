import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
base_dir = r"d:\!!MYDOCUMENTS2026\!!!SISTEKIN2026\!!BRAINSTORMING_KURIKULUM2026\KURIKULUM2026_REVISI"

# Check 002-RPS-MSN-315 Dasar Pemrograman Komputer.docx
doc_path = os.path.join(base_dir, "002-RPS-MSN-315 Dasar Pemrograman Komputer.docx")
if os.path.exists(doc_path):
    doc = docx.Document(doc_path)
    print("DOCX: 002-RPS-MSN-315 Dasar Pemrograman Komputer.docx")
    print(f"Total tables: {len(doc.tables)}")
    for i, t in enumerate(doc.tables):
        print(f"Table {i+1}: {len(t.rows)} rows x {len(t.columns)} cols")
        if len(t.rows) > 0 and len(t.columns) > 0:
            print("  First row sample:", [c.text.strip().replace('\n', ' ')[:40] for c in t.rows[0].cells[:4]])

# Also check 0003 prompt file
prompt_path = os.path.join(base_dir, "0003_PROMPT - BUAT RPS OBE With AI.pdf")
if os.path.exists(prompt_path):
    import fitz
    pdoc = fitz.open(prompt_path)
    print(f"\n0003_PROMPT: pages={len(pdoc)}")
    for i in range(len(pdoc)):
        txt = pdoc[i].get_text("text")
        prompts = [line for line in txt.split("\n") if "Prompt " in line]
        if prompts:
            print(f"  Page {i+1} prompts: {prompts}")
