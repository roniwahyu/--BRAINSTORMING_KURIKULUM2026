# -*- coding: utf-8 -*-
"""Table 18 consistency fix:
Append the 2 missing 0-SKS MKWU courses (MKU-406 Agama II @ Sem4, MKU-508 Kewirausahaan II @ Sem5)
so the MK list totals 55 MK, matching Table 19 rekap & Dok 005 consensus.

Because Table 18 currently renders 53 rows (all SKS>0) sorted by No with 'No' column
renumbered 1..53, the cleanest deterministic fix is:
  1. Read all data rows into records.
  2. Insert the 2 missing courses at the correct semester position.
  3. Rebuild the body by cloning the header-row template (proven duplicate_row pattern),
     wipe text, and fill every record + renumber.

The reconstruction approach (delete all body rows, append fresh clones) avoids the
stale-object traps of surgical in-place insertion.
"""
import copy
from docx import Document
from docx.oxml.ns import qn

SRC = "D:/!!MYDOCUMENTS2026/!!!SISTEKIN2026/!!BRAINSTORMING_KURIKULUM2026/KURIKULUM2026_REVISI/BUKU_KPT_SISTEKIN_2026_FINAL.docx"

def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold

def rv(r):
    return [c.text.strip() for c in r.cells]

doc = Document(SRC)
t = doc.tables[18]

# --- 1. Collect current data records (body rows that carry MK info) ---
records = []
for r in t.rows:
    joined = " ".join(rv(r)).upper()
    if "KODE MK" in joined and "NAMA MATA KULIAH" in joined:
        continue  # header
    if any(k in joined for k in ("TOTAL", "JUMLAH", "SUBTOTAL")):
        continue  # summary rows (if any). We keep header & summary separately.
    cells = rv(r)
    if len(cells) < 5 or not cells[1]:
        continue
    try:
        sem = int(cells[4])
    except ValueError:
        continue
    records.append({
        "kode": cells[1], "nama": cells[2], "sks": cells[3], "sem": sem,
        "jenis": cells[5] if len(cells) > 5 else "",
    })

# --- 2. Insert the 2 missing 0-SKS courses by semester position ---
INSERT = [
    {"kode": "MKU-406", "nama": "Agama II", "sks": "0", "sem": 4, "jenis": "MKWU"},
    {"kode": "MKU-508", "nama": "Kewirausahaan II", "sks": "0", "sem": 5, "jenis": "MKWU"},
]
for ins in INSERT:
    at = len(records)
    for idx, rec in enumerate(records):
        if rec["sem"] > ins["sem"]:
            at = idx
            break
    records.insert(at, dict(ins))

# --- 3. Renumber ---
for n, rec in enumerate(records, start=1):
    rec["no"] = str(n)

# --- 4. Identify header row + summary rows to preserve ---
header_tr = None
summary_trs = []
for r in t.rows:
    joined = " ".join(rv(r)).upper()
    if "KODE MK" in joined and "NAMA MATA KULIAH" in joined:
        header_tr = r._tr
    elif any(k in joined for k in ("TOTAL", "JUMLAH", "SUBTOTAL")):
        summary_trs.append(r._tr)

if header_tr is None:
    header_tr = t.rows[0]._tr

# --- 5. Drop ALL rows from the table, then rebuild: header + N clones + summaries ---
tbl_el = t._tbl
for child in list(tbl_el):
    if child.tag == qn('w:tr'):
        tbl_el.remove(child)
tbl_el.append(header_tr)  # header first

# Build a blank template clone (strip run text) to fill per record
def blank_clone(src_tr):
    tr = copy.deepcopy(src_tr)
    for tc in tr.findall(qn('w:tc')):
        for p in tc.findall(qn('w:p')):
            for r_ in p.findall(qn('w:r')):
                p.remove(r_)
    return tr

blank_tmpl = blank_clone(header_tr)

for rec in records:
    tr = copy.deepcopy(blank_tmpl)
    tbl_el.append(tr)
for s in summary_trs:
    tbl_el.append(s)

# --- 6. Fill the body cells (save rebuilt structure first, then reload+filling) ---
doc.save(SRC)  # persist the new row structure
doc2 = Document(SRC)  # reload fully rebuilt file then fill
t2 = doc2.tables[18]
# data rows = rows after header, before summaries -> first len(records) rows after index 0
tr_elms = t2._tbl.findall(qn('w:tr'))
# Data rows are indices 1..1+len(records) (header at 0, records after, summaries last)
from docx.table import _Row
data_els = tr_elms[1:1+len(records)]
assert len(data_els) == len(records), f"expected {len(records)} data rows, got {len(data_els)}"
for i, te in enumerate(data_els):
    rec = records[i]
    roww = _Row(te, t2)
    cells = roww.cells
    vals = [rec["no"], rec["kode"], rec["nama"], rec["sks"], str(rec["sem"]), rec["jenis"]]
    for j, v in enumerate(vals):
        if j < len(cells):
            set_cell(cells[j], v)

doc2.save(SRC)
print(f"SAVED. Table 18 rebuilt with {len(records)} MK rows (was 53).")
